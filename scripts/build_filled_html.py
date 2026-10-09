# -*- coding: utf-8 -*-
import openpyxl, html, re
from pypinyin import lazy_pinyin

WB = openpyxl.load_workbook("data/filled_from_gsheet.xlsx", data_only=True)

# ---------- parse 依桌次 ----------
ws = WB["依桌次"]
tables = []  # {name, loc, seats:[(seat, name, rel, note)]}
cur = None
hdr_re = re.compile(r"^　?(.+?)　—　(.+)$")
for r in range(1, ws.max_row + 1):
    a = ws.cell(r,1).value; b = ws.cell(r,2).value
    c = ws.cell(r,3).value; d = ws.cell(r,4).value
    if a is None and b is None and c is None and d is None:
        continue
    sa = str(a).strip() if a is not None else ""
    # table header row (merged into col A, contains —)
    m = hdr_re.match(sa)
    if m:
        name = m.group(1).replace("（主桌）","").strip()
        loc  = m.group(2).strip()
        cur = {"name": name, "loc": loc, "seats": []}
        tables.append(cur)
        continue
    if sa == "座位":   # sub-header
        continue
    if sa.startswith("　本桌") or sa.startswith("本桌"):
        continue
    if "依桌次" in sa and "排桌" in sa:
        continue
    # seat row: col A numeric
    if cur is not None and isinstance(a,(int,float)):
        nm = (str(b).strip() if b not in (None,"") else "")
        rel= (str(c).strip() if c not in (None,"") else "")
        nt = (str(d).strip() if d not in (None,"") else "")
        cur["seats"].append((int(a), nm, rel, nt))

# ---------- parse 賓客總表 ----------
mw = WB["賓客總表"]
guests = []  # dict
TABLE_ORDER = [t["name"] for t in tables]
for r in range(4, mw.max_row+1):
    name = mw.cell(r,2).value
    tb   = mw.cell(r,3).value
    rel  = mw.cell(r,4).value
    cake = mw.cell(r,7).value
    pax  = mw.cell(r,8).value
    note = mw.cell(r,10).value
    if name in (None,"") or (isinstance(name,str) and name.startswith("（範例")):
        continue
    if isinstance(name,str) and ("：" in name or "總數" in name):
        continue
    guests.append({
        "name": str(name).strip(),
        "table": str(tb).strip() if tb else "",
        "rel": str(rel).strip() if rel else "",
        "cake": str(cake).strip() if cake else "",
        "pax": int(pax) if isinstance(pax,(int,float)) else (1 if tb else None),
        "note": str(note).strip() if note else "",
    })

# ---------- stats ----------
def seat_used(t):
    return sum(1 for (s,n,rel,nt) in t["seats"] if n or rel)
total_units = len(guests)
# headcount estimate: sum 報名人數 (blank treated as 1) over 總表
def pax_of(g): return g["pax"] if g["pax"] else 1
total_head = sum(pax_of(g) for g in guests)
# working staff on 睡美人 (in 依桌次 but not 總表)
STAFF_LABELS = {"主持人","音控","新秘","新秘助手"}
staff = []
for t in tables:
    for (s,n,rel,nt) in t["seats"]:
        if rel in STAFF_LABELS:
            staff.append((t["name"], n, rel))

# ---------- name lookup (sorted by pinyin) ----------
def sort_key(g):
    nm = g["name"]
    # use first token before / for pinyin, fall back to the whole
    base = nm.split("/")[0]
    py = lazy_pinyin(base)
    # if ascii (English), lowercase
    first = py[0] if py else nm
    return (first.lower(), nm)
lookup = sorted(guests, key=sort_key)

# =====================================================================
# build HTML
# =====================================================================
def esc(x): return html.escape(str(x)) if x is not None else ""

# ---- Section 1 cards ----
def card(t):
    used = seat_used(t)
    is_head = (t["name"] == "王位")
    seats = t["seats"]
    # trim trailing fully-empty, keep up to 2 spare blanks
    last = 0
    for i,(s,n,rel,nt) in enumerate(seats):
        if n or rel: last = i+1
    cap = min(len(seats), last + 2)
    rows = []
    for (s,n,rel,nt) in seats[:cap]:
        blankname = "" if n else '<span class="ph">—</span>'
        rows.append(
            f'<tr><td class="num">{s}</td>'
            f'<td class="nm">{esc(n) or blankname}</td>'
            f'<td class="rel">{esc(rel)}</td>'
            f'<td class="note">{esc(nt)}</td></tr>')
    rows_html = "\n".join(rows)
    return f'''
<section class="card {'main' if is_head else ''}">
  <div class="card-head">
    <div class="tname">{esc(t["name"])}{' <span class="badge">主桌</span>' if is_head else ''}</div>
    <div class="tloc">{esc(t["loc"])}</div>
    <div class="tcount">本桌 {used} 位</div>
  </div>
  <table class="seat">
    <thead><tr><th class="num">座位</th><th class="nm">姓名</th>
    <th class="rel">關係／稱謂</th><th class="note">備註</th></tr></thead>
    <tbody>
{rows_html}
    </tbody>
  </table>
</section>'''
cards_html = "\n".join(card(t) for t in tables)

# ---- Section 2 禮金桌核對表 (by table order, named units) ----
gift_rows = []
torder = {n:i for i,n in enumerate(TABLE_ORDER)}
gsorted = sorted(guests, key=lambda g:(torder.get(g["table"],99), ))
for g in gsorted:
    pax = g["pax"] if g["pax"] else ""
    gift_rows.append(
        f'<tr><td class="g-nm">{esc(g["name"])}</td>'
        f'<td class="g-rel">{esc(g["rel"])}</td>'
        f'<td class="g-tb">{esc(g["table"])}</td>'
        f'<td class="g-pax">{pax}</td>'
        f'<td class="g-amt"></td><td class="g-ck"></td></tr>')
# add a few blank rows for walk-ins
for _ in range(6):
    gift_rows.append('<tr><td class="g-nm"></td><td class="g-rel"></td><td class="g-tb"></td>'
                     '<td class="g-pax"></td><td class="g-amt"></td><td class="g-ck"></td></tr>')
gift_html = "\n".join(gift_rows)

# ---- Section 3 姓名查詢 (two columns) ----
lk_rows = []
for g in lookup:
    lk_rows.append(
        f'<tr><td class="l-nm">{esc(g["name"])}</td>'
        f'<td class="l-tb">{esc(g["table"])}</td></tr>')
# split into 2 balanced columns
half = (len(lk_rows)+1)//2
col1 = "\n".join(lk_rows[:half]); col2 = "\n".join(lk_rows[half:])

staff_note = ""
if staff:
    items = "、".join(f'{esc(n) or esc(rel)}（{esc(rel)}）' for (tb,n,rel) in staff)
    stbl = staff[0][0] if staff else ""
    staff_note = f'工作人員（{items}）安排於「{esc(stbl)}」桌。'

HTML = f'''<!doctype html>
<html lang="zh-Hant"><head>
<meta charset="utf-8"><meta name="viewport" content="width=device-width, initial-scale=1">
<title>訂婚賓客對照表</title>
<meta name="description" content="Barry ♥ Winnie 訂婚宴 接待・禮金桌・帶位用賓客對照表（已填名單），宮廷紅金可列印版。">
<link rel="preconnect" href="https://fonts.googleapis.com">
<link rel="preconnect" href="https://fonts.gstatic.com" crossorigin>
<link href="https://fonts.googleapis.com/css2?family=Cormorant+Garamond:wght@500;600;700&family=Noto+Serif+TC:wght@500;600;700;900&display=swap" rel="stylesheet">
<style>
  :root{{--red:#8B1A1A;--wine:#6E1414;--gold:#C9A24B;--gold-d:#A8822F;
    --gold-l:#F3E7C8;--cream:#FBF6EA;--ink:#2B1D12;--line:#D8C9A6;}}
  *{{box-sizing:border-box;}}
  html,body{{margin:0;padding:0;background:#efe6d2;color:var(--ink);font-family:"Noto Serif TC",serif;}}
  .wrap{{max-width:960px;margin:0 auto;padding:24px 16px 60px;}}
  .cover{{text-align:center;padding:34px 20px 26px;margin-bottom:22px;
    background:linear-gradient(160deg,#7d1616,#8B1A1A 55%,#6E1414);
    border:2px solid var(--gold);border-radius:14px;color:#fff;box-shadow:0 10px 30px rgba(0,0,0,.18);}}
  .cover .ey{{font-family:"Cormorant Garamond",serif;letter-spacing:.42em;font-size:14px;color:var(--gold-l);text-transform:uppercase;}}
  .cover h1{{font-weight:900;font-size:33px;margin:10px 0 6px;text-shadow:0 2px 10px rgba(0,0,0,.25);}}
  .cover .names{{font-family:"Cormorant Garamond",serif;font-size:22px;letter-spacing:.14em;color:var(--gold-l);}}
  .cover .meta{{margin-top:12px;font-size:13px;color:#f0ddc0;letter-spacing:.2em;}}
  .cover .rule{{width:70px;height:2px;background:var(--gold);margin:14px auto 0;}}
  .stats{{display:flex;gap:10px;flex-wrap:wrap;justify-content:center;margin:18px 0 26px;}}
  .stat{{background:var(--cream);border:1px solid var(--line);border-radius:10px;padding:10px 18px;text-align:center;min-width:120px;}}
  .stat b{{display:block;font-size:24px;color:var(--red);font-weight:900;}}
  .stat span{{font-size:12px;color:var(--wine);}}
  .staff{{background:var(--gold-l);border:1px solid var(--gold);border-radius:10px;
    padding:10px 16px;margin:0 0 24px;font-size:13px;color:var(--wine);text-align:center;}}
  .sec-h{{font-weight:900;color:var(--red);font-size:22px;margin:30px 0 12px;padding-bottom:6px;border-bottom:2px solid var(--gold);}}
  .sec-h .en{{font-family:"Cormorant Garamond",serif;color:var(--gold-d);font-size:15px;letter-spacing:.2em;margin-left:10px;}}
  .card{{background:#fff;border:1.5px solid var(--gold);border-radius:12px;padding:12px 16px 14px;margin:0 0 16px;box-shadow:0 4px 14px rgba(0,0,0,.06);break-inside:avoid;}}
  .card.main{{border-width:2.5px;background:#fffdf6;}}
  .card-head{{display:flex;align-items:baseline;gap:14px;flex-wrap:wrap;padding-bottom:7px;margin-bottom:7px;border-bottom:1px dashed var(--line);}}
  .tname{{font-weight:900;font-size:21px;color:var(--red);}}
  .tname .badge{{font-size:12px;background:var(--gold);color:#4a2a08;padding:2px 8px;border-radius:10px;vertical-align:middle;margin-left:6px;}}
  .tloc{{font-size:13px;color:var(--wine);}}
  .tcount{{margin-left:auto;font-size:14px;color:var(--ink);font-weight:700;}}
  table{{width:100%;border-collapse:collapse;}}
  .seat th,.seat td{{border:1px solid var(--line);padding:6px 8px;font-size:14px;}}
  .seat th{{background:var(--gold-l);color:var(--wine);font-weight:700;}}
  .seat .num{{width:48px;text-align:center;color:#9a7b3f;}}
  .seat .nm{{width:24%;font-weight:600;}} .seat .rel{{width:30%;}}
  .seat .ph{{color:#c4b48f;}}
  .seat tbody td{{height:26px;}}
  .sheet{{background:#fff;border:1.5px solid var(--gold);border-radius:12px;padding:14px 16px;margin:0 0 18px;break-inside:avoid;}}
  .sheet-title{{font-weight:900;color:var(--red);font-size:19px;margin-bottom:6px;display:flex;align-items:baseline;gap:12px;}}
  .sheet-title .pg{{font-size:13px;color:var(--gold-d);font-weight:600;}}
  .hint{{font-size:12px;color:var(--wine);margin:0 0 10px;font-style:italic;}}
  .gift th,.gift td{{border:1px solid var(--line);padding:6px 8px;font-size:13.5px;}}
  .gift th{{background:var(--gold-l);color:var(--wine);}}
  .gift .g-nm{{width:20%;font-weight:600;}} .gift .g-rel{{width:24%;}}
  .gift .g-tb{{width:14%;text-align:center;}} .gift .g-pax{{width:9%;text-align:center;}}
  .gift .g-amt{{width:20%;}} .gift .g-ck{{width:9%;text-align:center;}}
  .gift tbody td{{height:28px;}}
  .lk-wrap{{display:flex;gap:14px;}}
  .lookup{{flex:1;}}
  .lookup th,.lookup td{{border:1px solid var(--line);padding:5px 8px;font-size:13.5px;}}
  .lookup th{{background:var(--gold-l);color:var(--wine);}}
  .lookup .l-nm{{font-weight:600;}} .lookup .l-tb{{width:42%;}}
  .foot{{text-align:center;color:var(--wine);font-size:12px;margin-top:30px;letter-spacing:.12em;}}
  .foot .glyph{{color:var(--gold-d);margin:0 6px;}}
  .top-nav{{display:flex;justify-content:space-between;align-items:center;margin-bottom:14px;gap:8px;flex-wrap:wrap;}}
  .nav-btn{{display:inline-flex;align-items:center;gap:6px;font-size:13px;color:var(--wine);text-decoration:none;background:#fff;border:1px solid var(--line);padding:6px 14px;border-radius:999px;font-weight:600;box-shadow:0 2px 6px rgba(0,0,0,0.04);transition:all .15s;}}
  .nav-btn.gsheet{{background:var(--cream);border-color:var(--gold);font-weight:700;}}
  .nav-btn:hover{{background:var(--gold-l);border-color:var(--gold);}}
  @media print{{
    html,body{{background:#fff;}} .wrap{{max-width:none;padding:0;}}
    .top-nav{{display:none !important;}}
    .cover,.card,.sheet,.stat,.staff,th,td{{-webkit-print-color-adjust:exact;print-color-adjust:exact;}}
    .card,.sheet{{box-shadow:none;page-break-inside:avoid;}}
    .cover{{box-shadow:none;}}
    @page{{size:A4;margin:12mm;}}
  }}
</style></head>
<body><div class="wrap">
  <div class="top-nav">
    <a href="index.html" class="nav-btn">‹ 返回工作台首頁</a>
    <a href="https://docs.google.com/spreadsheets/d/1TasCO9KSlpflyi-poXNur0p5EYWLFvtHGZI3R7eKsYg/edit?gid=848591360#gid=848591360" target="_blank" rel="noopener" class="nav-btn gsheet">📊 Google 試算表原始母檔 ↗</a>
  </div>

  <div class="cover">
    <div class="ey">Seating &amp; Reception Guide</div>
    <h1>訂婚宴 ・ 賓客對照表</h1>
    <div class="names">BARRY &amp; WINNIE</div>
    <div class="rule"></div>
    <div class="meta">幸福登機門　GATE 2F　・　FLIGHT BW1011　・　2026.10.11　台南</div>
  </div>

  <div class="stats">
    <div class="stat"><b>13</b><span>桌次</span></div>
    <div class="stat"><b>{total_units}</b><span>名單組數</span></div>
    <div class="stat"><b>約 {total_head}</b><span>預估入席人數</span></div>
  </div>
  {f'<div class="staff">{staff_note}</div>' if staff_note else ''}

  <div class="sec-h">① 依桌次 <span class="en">BY TABLE ・ 帶位・桌卡</span></div>
  {cards_html}

  <div class="sec-h">② 禮金桌核對表 <span class="en">GIFT DESK</span></div>
  <div class="sheet">
    <div class="sheet-title">禮金桌核對表</div>
    <p class="hint">依桌次排序。「組數」為報名人數，收禮金時可核對同行人數；金額與打勾欄當天填寫。末尾留白供臨時賓客。</p>
    <table class="gift">
      <thead><tr><th class="g-nm">姓名</th><th class="g-rel">關係／稱謂</th><th class="g-tb">桌次</th>
      <th class="g-pax">組數</th><th class="g-amt">紅包金額</th><th class="g-ck">✓</th></tr></thead>
      <tbody>
{gift_html}
      </tbody>
    </table>
  </div>

  <div class="sec-h">③ 賓客查詢表 <span class="en">NAME → TABLE ・ 依姓名</span></div>
  <div class="sheet">
    <div class="sheet-title">賓客查詢表<span class="pg">依姓名排序　共 {total_units} 筆</span></div>
    <p class="hint">賓客報名字即可查到桌次。同行親友（如「○○老公／女友／小孩」）與主要賓客同桌，請一起帶位。</p>
    <div class="lk-wrap">
      <table class="lookup"><thead><tr><th class="l-nm">姓名</th><th class="l-tb">桌次</th></tr></thead>
        <tbody>{col1}</tbody></table>
      <table class="lookup"><thead><tr><th class="l-nm">姓名</th><th class="l-tb">桌次</th></tr></thead>
        <tbody>{col2}</tbody></table>
    </div>
  </div>

  <div class="foot">BARRY <span class="glyph">&#10022;</span> WINNIE<span class="glyph">&#10022;</span>幸福登機門　TAINAN 10&middot;11</div>
</div></body></html>'''

with open("outputs/html/訂婚賓客對照表_已填_可列印.html","w",encoding="utf-8") as f:
    f.write(HTML)
with open("docs/guests.html","w",encoding="utf-8") as f:
    f.write(HTML)

# print a data-health report to stdout
print("TABLES:", len(tables))
print("GUEST UNITS:", total_units, " EST HEAD:", total_head)
print("STAFF:", staff)
print("\nper-table seats used (依桌次) vs 總表組數:")
from collections import Counter
cnt_master = Counter(g["table"] for g in guests)
for t in tables:
    print(f'  {t["name"]:<6} seats_used={seat_used(t):>2}  master_units={cnt_master.get(t["name"],0):>2}')
EOF_MARKER = None
